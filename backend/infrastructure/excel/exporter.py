"""Minimal order workbook; the partner-specific 1C template is not specified yet."""

from collections.abc import Mapping
from io import BytesIO
from uuid import UUID

from openpyxl import Workbook

from backend.domain.entities.catalog import Supplier, Warehouse
from backend.domain.entities.product import Product
from backend.domain.entities.purchase_order import PurchaseOrder


class XlsxOrderExporter:
    def render(self, order: PurchaseOrder, *, supplier: Supplier, warehouse: Warehouse,
               products: Mapping[UUID, Product]) -> bytes:
        workbook = Workbook()
        try:
            sheet = workbook.active
            sheet.title = "Order"
            rows = [
                ["order_number", order.order_number],
                ["order_id", str(order.id)],
                ["supplier_code", supplier.code], ["supplier_name", supplier.name],
                ["warehouse_code", warehouse.code], ["warehouse_name", warehouse.name],
                ["approved_at", order.approved_at.isoformat() if order.approved_at else ""],
                [],
                ["sku", "product_name", "unit", "recommended_quantity", "approved_quantity",
                 "unit_price", "total_amount", "recommendation_id"],
            ]
            for item in order.items:
                product = products[item.product_id]
                # Decimal text preserves all 18 digits; Excel numeric cells only preserve 15.
                rows.append([product.sku, product.name, product.unit,
                             format(item.recommended_quantity, ".4f"), format(item.approved_quantity, ".4f"),
                             format(item.unit_price, ".4f") if item.unit_price is not None else "",
                             format(item.total_amount, ".4f") if item.total_amount is not None else "",
                             str(item.recommendation_id)])
            for row in rows:
                sheet.append(row)
                for cell in sheet[sheet.max_row]:
                    if isinstance(cell.value, str):
                        # Catalog values such as '=...' must stay literal text, never formulas.
                        cell.data_type = "s"
            sheet.freeze_panes = "A10"
            sheet.auto_filter.ref = f"A9:H{sheet.max_row}"
            with BytesIO() as stream:
                workbook.save(stream)
                return stream.getvalue()
        finally:
            workbook.close()
