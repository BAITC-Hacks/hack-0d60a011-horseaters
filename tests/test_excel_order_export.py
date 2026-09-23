import unittest
from datetime import date, datetime, timezone
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook

from backend.application.dto.order import OrderExportRow
from backend.application.dto.order import OrderExportCommand
from backend.application.use_cases.export_order import ExportOrder, OrderNotExportableError
from backend.domain.entities.catalog import Supplier, SupplierProduct, Warehouse
from backend.domain.entities.product import Product
from backend.domain.entities.purchase_order import PurchaseOrder, PurchaseOrderItem
from backend.infrastructure.excel.exporter import (
    ORDER_EXPORT_COLUMNS,
    OpenpyxlOrderWorkbookExporter,
)


SAMPLE_ROW = OrderExportRow(
    order_number="PO-2026-0001",
    supplier_name="ИЭК",
    warehouse_name="Основной склад",
    sku="A-100",
    product_name="Автоматический выключатель",
    recommended_quantity=Decimal("120"),
    approved_quantity=Decimal("100"),
    unit_price=Decimal("245.50"),
    total_amount=Decimal("24550.00"),
    delivery_date=date(2026, 10, 7),
)


class ExcelOrderExportTests(unittest.TestCase):
    def assert_contract(self, workbook_bytes: bytes) -> None:
        workbook = load_workbook(BytesIO(workbook_bytes), data_only=True)
        self.addCleanup(workbook.close)
        sheet = workbook["Заказ"]

        self.assertEqual(tuple(cell.value for cell in sheet[1]), ORDER_EXPORT_COLUMNS)
        self.assertEqual(len(sheet.merged_cells.ranges), 0)
        for column in range(6, 10):
            self.assertEqual(sheet.cell(2, column).data_type, "n")
        self.assertIsInstance(sheet.cell(2, 10).value, date)
        self.assertEqual(sheet.cell(2, 10).number_format, "dd.mm.yyyy")

    def test_export_has_fixed_columns_and_native_excel_types(self):
        content = OpenpyxlOrderWorkbookExporter().render([SAMPLE_ROW])
        self.assert_contract(content)

    def test_committed_control_example_matches_contract(self):
        sample = Path(__file__).parents[1] / "docs" / "examples" / "order-export-1c.xlsx"
        self.assertTrue(sample.is_file(), "control XLSX example is missing")
        self.assert_contract(sample.read_bytes())


class _SingleValueRepository:
    def __init__(self, value):
        self.value = value

    def get(self, _identifier):
        return self.value

    def get_by_id(self, _identifier):
        return self.value


class _SupplierRepository(_SingleValueRepository):
    def __init__(self, supplier, terms):
        super().__init__(supplier)
        self.terms = terms

    def list_terms(self, *_args, **_kwargs):
        return [self.terms]


class _OrderRepository(_SingleValueRepository):
    def __init__(self, order):
        super().__init__(order)
        self.export = None

    def add_export(self, export):
        self.export = export
        return export


class _Uow:
    def __init__(self, *, order, supplier, warehouse, product, terms):
        self.orders = _OrderRepository(order)
        self.suppliers = _SupplierRepository(supplier, terms)
        self.warehouses = _SingleValueRepository(warehouse)
        self.products = _SingleValueRepository(product)
        self.committed = False

    def __enter__(self):
        return self

    def __exit__(self, *_args):
        return None

    def commit(self):
        self.committed = True


class _CapturingExporter:
    def __init__(self):
        self.rows = []

    def render(self, rows):
        self.rows = list(rows)
        return b"xlsx-content"


class ExportOrderUseCaseTests(unittest.TestCase):
    def setUp(self):
        self.user_id = uuid4()
        self.product = Product(id=uuid4(), sku="A-100", name="Автомат", unit="шт")
        self.supplier = Supplier(id=uuid4(), code="IEK", name="ИЭК")
        self.warehouse = Warehouse(id=uuid4(), code="MAIN", name="Основной")
        self.terms = SupplierProduct(
            id=uuid4(), supplier_id=self.supplier.id, product_id=self.product.id,
            moq=Decimal("10"), package_size=Decimal("5"), lead_time_days=14,
        )
        self.order = PurchaseOrder(
            order_number="PO/2026 0001", supplier_id=self.supplier.id,
            warehouse_id=self.warehouse.id, created_from_run_id=uuid4(),
            created_by=self.user_id,
            created_at=datetime(2026, 9, 20, tzinfo=timezone.utc),
        )
        self.order.add_item(PurchaseOrderItem(
            recommendation_id=uuid4(), product_id=self.product.id,
            recommended_quantity=Decimal("120"), approved_quantity=Decimal("100"),
            unit_price=Decimal("2.50"), total_amount=None,
        ))

    def uow(self):
        return _Uow(
            order=self.order, supplier=self.supplier, warehouse=self.warehouse,
            product=self.product, terms=self.terms,
        )

    def test_approved_order_is_rendered_and_audited_atomically(self):
        self.order.approve(
            approved_by=self.user_id,
            approved_at=datetime(2026, 9, 23, tzinfo=timezone.utc),
        )
        uow = self.uow()
        exporter = _CapturingExporter()
        result = ExportOrder(lambda: uow, exporter).execute(
            OrderExportCommand(order_id=self.order.id, user_id=self.user_id)
        )

        self.assertTrue(uow.committed)
        self.assertEqual(result.file_name, "order-PO-2026-0001.xlsx")
        self.assertEqual(result.checksum, sha256(b"xlsx-content").hexdigest())
        self.assertEqual(exporter.rows[0].delivery_date, date(2026, 10, 7))
        self.assertEqual(exporter.rows[0].total_amount, Decimal("250.00"))
        self.assertEqual(uow.orders.export.file_checksum, result.checksum)

    def test_draft_order_cannot_be_exported(self):
        with self.assertRaises(OrderNotExportableError):
            ExportOrder(lambda: self.uow(), _CapturingExporter()).execute(
                OrderExportCommand(order_id=self.order.id, user_id=self.user_id)
            )


if __name__ == "__main__":
    unittest.main()
